from flask import Blueprint, request, jsonify
from flask_jwt_extended import jwt_required
import os
import requests
import json
import io
import pypdf
import docx
from server.models import db, Quiz, Question
from server.middleware.auth import admin_required

quizzes_bp = Blueprint('quizzes', __name__)

@quizzes_bp.route('', methods=['GET'])
@jwt_required()
def get_all_quizzes():
    try:
        # Extract search, filter, and pagination params
        topic = request.args.get('topic', '').strip()
        difficulty = request.args.get('difficulty', '').strip()
        search = request.args.get('search', '').strip()
        
        page = request.args.get('page', type=int)
        limit = request.args.get('limit', type=int, default=10)

        query = Quiz.query

        if topic:
            query = query.filter_by(topic=topic)
        if difficulty:
            query = query.filter_by(difficulty=difficulty)
        if search:
            query = query.filter(Quiz.title.ilike(f'%{search}%'))

        # Check for paginated vs full request
        if page:
            pagination = query.order_by(Quiz.created_at.desc()).paginate(page=page, per_page=limit, error_out=False)
            quizzes_list = [quiz.to_dict() for quiz in pagination.items]
            return jsonify({
                'quizzes': quizzes_list,
                'total': pagination.total,
                'page': pagination.page,
                'pages': pagination.pages,
                'has_next': pagination.has_next,
                'has_prev': pagination.has_prev
            }), 200
        else:
            quizzes = query.order_by(Quiz.created_at.desc()).all()
            return jsonify({'quizzes': [q.to_dict() for q in quizzes]}), 200
            
    except Exception as e:
        return jsonify({'error': f'Failed to retrieve quizzes: {str(e)}'}), 500

@quizzes_bp.route('/<int:quiz_id>', methods=['GET'])
@jwt_required()
def get_quiz_by_id(quiz_id):
    try:
        quiz = Quiz.query.get(quiz_id)
        if not quiz:
            return jsonify({'error': 'Quiz not found.'}), 404
        
        # Structure the payload with the questions
        questions = Question.query.filter_by(quiz_id=quiz.id).all()
        quiz_data = quiz.to_dict()
        quiz_data['questions'] = [q.to_dict() for q in questions]

        return jsonify({'quiz': quiz_data}), 200
    except Exception as e:
        return jsonify({'error': f'Failed to load quiz: {str(e)}'}), 500

@quizzes_bp.route('', methods=['POST'])
@admin_required()
def create_quiz():
    data = request.get_json() or {}
    title = data.get('title', '').strip()
    topic = data.get('topic', '').strip()
    difficulty = data.get('difficulty', '').strip()

    if not title or not topic or not difficulty:
        return jsonify({'error': 'Title, topic, and difficulty are required.'}), 400

    try:
        new_quiz = Quiz(title=title, topic=topic, difficulty=difficulty)
        db.session.add(new_quiz)
        db.session.commit()
        return jsonify({'message': 'Quiz created successfully.', 'quiz': new_quiz.to_dict()}), 201
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Failed to create quiz: {str(e)}'}), 500

@quizzes_bp.route('/<int:quiz_id>', methods=['PUT'])
@admin_required()
def update_quiz(quiz_id):
    try:
        quiz = Quiz.query.get(quiz_id)
        if not quiz:
            return jsonify({'error': 'Quiz not found.'}), 404

        data = request.get_json() or {}
        quiz.title = data.get('title', quiz.title).strip()
        quiz.topic = data.get('topic', quiz.topic).strip()
        quiz.difficulty = data.get('difficulty', quiz.difficulty).strip()

        db.session.commit()
        return jsonify({'message': 'Quiz updated successfully.', 'quiz': quiz.to_dict()}), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Failed to update quiz: {str(e)}'}), 500

@quizzes_bp.route('/<int:quiz_id>', methods=['DELETE'])
@admin_required()
def delete_quiz(quiz_id):
    try:
        quiz = Quiz.query.get(quiz_id)
        if not quiz:
            return jsonify({'error': 'Quiz not found.'}), 404

        db.session.delete(quiz)
        db.session.commit()
        return jsonify({'message': 'Quiz deleted successfully.'}), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Failed to delete quiz: {str(e)}'}), 500

def extract_json_from_text(text):
    text = text.strip()
    
    # Try parsing directly first
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
        
    # Find boundary indices of curly braces {...} or square brackets [...]
    first_brace = text.find('{')
    last_brace = text.rfind('}')
    first_bracket = text.find('[')
    last_bracket = text.rfind(']')
    
    start_idx = -1
    end_idx = -1
    
    if first_brace != -1 and last_brace != -1:
        if first_bracket != -1 and first_bracket < first_brace:
            start_idx = first_bracket
            end_idx = last_bracket
        else:
            start_idx = first_brace
            end_idx = last_brace
    elif first_bracket != -1 and last_bracket != -1:
        start_idx = first_bracket
        end_idx = last_bracket
        
    if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
        json_candidate = text[start_idx:end_idx+1]
        try:
            return json.loads(json_candidate)
        except json.JSONDecodeError as e:
            raise e
            
    raise json.JSONDecodeError("Could not find any JSON braces or brackets in output.", text, 0)

@quizzes_bp.route('/generate-from-syllabus', methods=['POST'])
@jwt_required()
def generate_from_syllabus():
    try:
        # Determine syllabus text
        syllabus_text = ''
        
        # Check if file was uploaded
        if 'file' in request.files:
            file = request.files['file']
            if file and file.filename != '':
                filename = file.filename.lower()
                file_bytes = file.read()
                
                if filename.endswith('.pdf'):
                    try:
                        pdf_file = io.BytesIO(file_bytes)
                        reader = pypdf.PdfReader(pdf_file)
                        text_content = []
                        for page in reader.pages:
                            page_text = page.extract_text()
                            if page_text:
                                text_content.append(page_text)
                        syllabus_text = "\n".join(text_content)
                    except Exception as pdf_err:
                        return jsonify({'error': f'Failed to parse PDF file: {str(pdf_err)}'}), 400
                elif filename.endswith('.docx'):
                    try:
                        docx_file = io.BytesIO(file_bytes)
                        doc = docx.Document(docx_file)
                        text_content = []
                        for paragraph in doc.paragraphs:
                            if paragraph.text:
                                text_content.append(paragraph.text)
                        for table in doc.tables:
                            for row in table.rows:
                                row_text = [cell.text for cell in row.cells if cell.text]
                                if row_text:
                                    text_content.append(" | ".join(row_text))
                        syllabus_text = "\n".join(text_content)
                    except Exception as docx_err:
                        return jsonify({'error': f'Failed to parse Word Document: {str(docx_err)}'}), 400
                else:
                    # Fallback to plain text decoding
                    syllabus_text = file_bytes.decode('utf-8', errors='ignore')
        
        # If syllabus text is still empty, look at form/json body
        if not syllabus_text:
            if request.is_json:
                data = request.get_json() or {}
                syllabus_text = data.get('syllabus_text', '')
            else:
                syllabus_text = request.form.get('syllabus_text', '')
        
        syllabus_text = syllabus_text.strip()
        if not syllabus_text:
            return jsonify({'error': 'Syllabus content or topic list is required. Please paste or upload a file.'}), 400

        # Determine total question count
        num_questions = 10
        if request.is_json:
            data = request.get_json() or {}
            num_questions = data.get('num_questions', 10)
        else:
            try:
                num_questions = int(request.form.get('num_questions', 10))
            except ValueError:
                num_questions = 10

        # Clamp between 5 and 30 questions
        num_questions = max(5, min(30, num_questions))

        # Calculate exact counts for easy/medium/hard (40-40-20 rule)
        count_easy = round(num_questions * 0.40)
        count_medium = round(num_questions * 0.40)
        count_hard = num_questions - count_easy - count_medium

        # API Key lookup
        # 1. Custom header X-Gemini-Key
        api_key = request.headers.get('X-Gemini-Key')
        if not api_key:
            if request.is_json:
                data = request.get_json() or {}
                api_key = data.get('api_key', '')
            else:
                api_key = request.form.get('api_key', '')
        
        # 2. Server environment
        if not api_key:
            api_key = os.environ.get('GEMINI_API_KEY')

        if not api_key:
            return jsonify({'error': 'Gemini API Key is missing. Please configure GEMINI_API_KEY in server/.env or input your key in the web interface.'}), 400

        # Prepare Gemma 4 payload
        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemma-4-31b-it:generateContent?key={api_key}"
        
        prompt = f"""
You are an expert technical interviewer and placement exam compiler.
Based on the following syllabus or list of topics, generate a highly personalized technical multiple-choice quiz:

---
{syllabus_text}
---

The quiz MUST contain exactly {num_questions} questions, broken down by difficulty as follows:
- Exactly {count_easy} Easy questions
- Exactly {count_medium} Medium questions
- Exactly {count_hard} Hard questions

Return the response ONLY as a JSON object, following this strict JSON schema:
{{
  "title": "A short, engaging title for the quiz, e.g., 'DBMS & SQL Mastery'",
  "topic": "A short topic label (max 15 chars, e.g. SQL, DBMS, OS, Java, C++)",
  "questions": [
    {{
      "question": "The question text",
      "difficulty": "Easy", // Must match the designated difficulty ("Easy", "Medium", or "Hard")
      "option_a": "Option A text",
      "option_b": "Option B text",
      "option_c": "Option C text",
      "option_d": "Option D text",
      "correct_answer": "A", // Must be "A", "B", "C", or "D"
      "explanation": "Detailed explanation of why the correct answer is right"
    }}
  ]
}}

Ensure all JSON rules are followed. Do not wrap the JSON output in markdown formatting or anything else, return pure JSON text. If you must use quotes in fields, escape them properly.
"""

        payload = {
            "contents": [{
                "parts": [{
                    "text": prompt
                }]
            }],
            "generationConfig": {
                "responseMimeType": "application/json"
            }
        }

        # Make HTTP Request
        print("[API] Attempting connection to Google Generative Language API (Gemma 4)...")
        try:
            response = requests.post(url, json=payload, timeout=120)
        except requests.exceptions.ConnectionError as conn_err:
            print(f"[API CONNECTION FAILED] Network connection could not be established: {str(conn_err)}")
            return jsonify({
                'error': 'Network connection failed: Unable to connect to Google API. Please ensure your machine is connected to the internet and DNS resolution is functioning properly.'
            }), 503
        except requests.exceptions.Timeout as timeout_err:
            print(f"[API CONNECTION TIMEOUT] The request to Google API timed out: {str(timeout_err)}")
            return jsonify({
                'error': 'API request timed out. Please check your network speed or try again.'
            }), 504
        except requests.exceptions.RequestException as req_err:
            print(f"[API ERROR] Request exception occurred: {str(req_err)}")
            return jsonify({
                'error': f'Failed to communicate with Google API: {str(req_err)}'
            }), 502

        if response.status_code != 200:
            print(f"[API HTTP ERROR] Connected, but API returned status {response.status_code}: {response.text}")
            return jsonify({'error': f'Gemini API request failed with status {response.status_code}: {response.text}'}), response.status_code

        print("[API CONNECTION SUCCESSFUL] Received valid response from Gemma 4 API.")

        res_data = response.json()
        try:
            raw_text = res_data['candidates'][0]['content']['parts'][0]['text'].strip()
        except (KeyError, IndexError):
            print(f"[API ERROR] Failed to extract text from Gemma response payload. Payload: {res_data}")
            return jsonify({'error': 'Failed to extract text from Gemma response payload.'}), 502

        try:
            quiz_json = extract_json_from_text(raw_text)
        except json.JSONDecodeError as e:
            print(f"[API PARSE ERROR] Failed to parse generated quiz JSON. Clean response was: {raw_text}. Error: {str(e)}")
            return jsonify({
                'error': f'Failed to parse generated quiz JSON. Error: {str(e)}',
                'raw_response': raw_text
            }), 502

        # Insert new Quiz
        quiz_title = quiz_json.get('title', 'Personalized Syllabus Quiz').strip()
        quiz_topic = quiz_json.get('topic', 'Syllabus').strip()
        quiz_difficulty = 'Personalized'

        new_quiz = Quiz(title=quiz_title, topic=quiz_topic, difficulty=quiz_difficulty)
        db.session.add(new_quiz)
        db.session.flush()  # Acquire ID before inserting questions

        questions_list = quiz_json.get('questions', [])
        if not questions_list:
            return jsonify({'error': 'AI generated an empty set of questions.'}), 502

        created_questions = []
        for idx, q_data in enumerate(questions_list):
            correct = q_data.get('correct_answer', 'A').strip().upper()
            if correct not in ['A', 'B', 'C', 'D']:
                correct = 'A'

            new_q = Question(
                quiz_id=new_quiz.id,
                question=q_data.get('question', f'Question {idx + 1}').strip(),
                option_a=q_data.get('option_a', 'Option A').strip(),
                option_b=q_data.get('option_b', 'Option B').strip(),
                option_c=q_data.get('option_c', 'Option C').strip(),
                option_d=q_data.get('option_d', 'Option D').strip(),
                correct_answer=correct,
                explanation=q_data.get('explanation', '').strip()
            )
            db.session.add(new_q)
            created_questions.append(new_q)

        db.session.commit()

        return jsonify({
            'message': 'Quiz generated successfully via AI!',
            'quiz_id': new_quiz.id,
            'title': new_quiz.title,
            'topic': new_quiz.topic,
            'question_count': len(created_questions)
        }), 201

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Failed to generate syllabus quiz: {str(e)}'}), 500

